"""Tests for PIN/signal Excel export."""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / ".github" / "skills" / "autosar-hw-code-traceability" / "scripts" / "export_pin_excel.py"


class PinExcelExportTests(unittest.TestCase):
    def test_exports_mcu_and_expander_pin_rows(self):
        inventory = {
            "components": [
                {"ref": "U1", "device_type": "MCU", "part_number": "MCU-PN", "package": "LQFP", "confidence": "high"},
                {"ref": "U2", "device_type": "IO expander", "part_number": "EXP-PN", "package": "QFN", "confidence": "medium"}
            ],
            "pins": [
                {"component_ref": "U1", "pin_number": "42", "pin_name_as_drawn": "PTA3", "datasheet_function": "GPIO", "net_id": "N1", "confidence": "high"},
                {"component_ref": "U2", "pin_number": "7", "pin_name_as_drawn": "IO0", "datasheet_function": "Expander GPIO", "net_id": "N1", "confidence": "medium"}
            ],
            "nets": [{"net_id": "N1", "name_as_drawn": "SENSOR_EN", "electrical_domain": "3V3", "signal_direction": "output", "endpoints": ["U1.42", "U2.7"]}]
        }
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "inventory.json"
            output = root / "pin_signals.xlsx"
            source.write_text(json.dumps(inventory), encoding="utf-8")
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--out", str(output)],
                capture_output=True, text=True, check=False
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertTrue(output.is_file())
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                ws = workbook["PIN信号清单"]
                self.assertEqual(ws.max_row, 3)
                headers = [cell.value for cell in ws[1]]
                row_mcu = [cell.value for cell in ws[2]]
                row_expander = [cell.value for cell in ws[3]]
                self.assertEqual(row_mcu[headers.index("器件类别")], "MCU")
                self.assertEqual(row_expander[headers.index("器件类别")], "扩展芯片")
                self.assertEqual(row_mcu[headers.index("网络名（原图）")], "SENSOR_EN")
                self.assertEqual(row_expander[headers.index("引脚号")], "7")
                self.assertIn("网络与跨页", workbook.sheetnames)
                self.assertIn("器件清单", workbook.sheetnames)
            finally:
                workbook.close()

    def test_prefers_explicit_function_owner_over_device_class_inference(self):
        inventory = {
            "components": [
                {"ref": "U1", "device_type": "MCU", "part_number": "MCU-PN", "package": "LQFP"}
            ],
            "pins": [
                {"component_ref": "U1", "pin_number": "1", "pin_name": "PTA0",
                 "net_id": "N1", "function_owner": "外部安全监控器"}
            ],
            "nets": [{"net_id": "N1", "name_as_drawn": "SAFE_MON"}]
        }
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "inventory.json"
            output = root / "pin_signals.xlsx"
            source.write_text(json.dumps(inventory), encoding="utf-8")
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--out", str(output)],
                capture_output=True, text=True, check=False
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                ws = workbook["PIN信号清单"]
                headers = [cell.value for cell in ws[1]]
                row = [cell.value for cell in ws[2]]
                self.assertEqual(row[headers.index("功能归属（MCU/扩展芯片/其他）")], "外部安全监控器")
            finally:
                workbook.close()

    def test_exports_chip_identity_review_without_auto_approval(self):
        inventory = {
            "components": [{
                "ref": "U1", "device_type": "MCU", "part_number": "BOM-PN",
                "chip_identity_review": {
                    "decision": "unconfirmed", "human_confirmation_required": True,
                    "checks": [{"part_number": {
                        "manifest_value": "BOM-PN",
                        "manual_candidates": [{"value": "MANUAL-PN", "pdf_page": 3, "evidence_excerpt": "Ordering Part Number: MANUAL-PN"}],
                        "status": "candidate_mismatch_review_required"
                    }}],
                    "review_gaps": [{"field": "package", "reason": "No manual candidate found"}]
                }
            }],
            "pins": [], "nets": []
        }
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, output = root / "inventory.json", root / "out.xlsx"
            source.write_text(json.dumps(inventory), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out", str(output)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                ws = workbook["器件清单"]
                headers = [cell.value for cell in ws[1]]
                row = [cell.value for cell in ws[2]]
                self.assertIn("未确认；人工复核必需", row[headers.index("手册身份核对状态")])
                self.assertIn("手册型号=MANUAL-PN (PDF p.3)", row[headers.index("手册候选型号/封装/修订")])
                self.assertIn("缺口:package", row[headers.index("身份核对证据")])
                self.assertEqual(row[headers.index("器件型号/料号")], "BOM-PN")
            finally:
                workbook.close()

    def test_exports_human_review_fields_without_replacing_source_state(self):
        inventory = {
            "components": [{"ref": "U1", "device_type": "MCU", "part_number": "PN-1"}],
            "pins": [{
                "component_ref": "U1", "pin_number": "4", "net_id": "N1",
                "review_state": "needs_review",
                "human_review": {"state": "verified", "entries": [{
                    "state": "verified", "reviewer": "reviewer-a", "reviewed_at": "2026-10-19",
                    "note": "Compared against controlled schematic",
                    "evidence": [{"filename": "schematic.pdf", "page": 2, "observation": "Pin label visible"}]
                }]}
            }],
            "nets": [{"net_id": "N1", "name_as_drawn": "IGN_SENSE"}]
        }
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, output = root / "inventory.json", root / "out.xlsx"
            source.write_text(json.dumps(inventory), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), "--out", str(output)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                ws = workbook["PIN信号清单"]
                headers = [cell.value for cell in ws[1]]
                row = [cell.value for cell in ws[2]]
                self.assertEqual(row[headers.index("审核状态")], "needs_review")
                self.assertIn("verified", row[headers.index("人工审核状态")])
                self.assertIn("reviewer-a", row[headers.index("审核人/时间")])
                self.assertIn("schematic.pdf / 2", row[headers.index("审核备注/证据")])
            finally:
                workbook.close()

    def test_rejects_invalid_inventory_arrays(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "inventory.json"
            output = root / "pin_signals.xlsx"
            source.write_text(json.dumps({"components": {}, "pins": [], "nets": []}), encoding="utf-8")
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--out", str(output)],
                capture_output=True, text=True, check=False
            )
            self.assertEqual(result.returncode, 2)
            self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
