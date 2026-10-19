#!/usr/bin/env python3
"""Export evidence-backed hardware inventory JSON to a multi-sheet Excel workbook."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


PIN_HEADERS = [
    "信号名称", "网络名（原图）", "网络ID", "器件类别", "器件位号",
    "器件型号/料号", "封装", "引脚号", "引脚名称（原图）", "引脚功能/复用功能",
    "功能归属（MCU/扩展芯片/其他）", "信号方向", "电气域/电压",
    "连接对象/端点", "原理图页码", "原理图页/Sheet", "数据手册引用",
    "证据说明", "置信度", "审核状态"
]
COMPONENT_HEADERS = [
    "器件位号", "器件类别", "器件型号/料号", "封装", "参数/数值",
    "装配状态", "功能说明（假设/已确认）", "置信度", "审核状态", "证据"
]
NET_HEADERS = [
    "网络ID", "网络名（原图）", "规范化网络名", "电气域/电压", "信号方向",
    "端点", "跨页源Sheet", "跨页源页码", "跨页源标签", "跨页目标Sheet",
    "跨页目标页码", "跨页目标标签", "跨页状态", "证据/原因", "置信度", "审核状态"
]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _evidence(items: Any) -> str:
    if not isinstance(items, list):
        return _text(items)
    parts = []
    for item in items:
        if isinstance(item, dict):
            loc = " / ".join(
                _text(item.get(k)) for k in ("document_id", "filename", "page", "sheet", "location", "xpath", "source")
                if item.get(k) not in (None, "")
            )
            observation = _text(item.get("observation") or item.get("quote") or item.get("note"))
            parts.append(": ".join(x for x in (loc, observation) if x))
        else:
            parts.append(_text(item))
    return " | ".join(x for x in parts if x)


def _device_class(component: dict[str, Any]) -> str:
    raw = _text(component.get("device_type") or component.get("component_type") or component.get("kind") or component.get("category")).strip()
    folded = raw.casefold().replace("_", " ").replace("-", " ")
    if any(token in folded for token in ("mcu", "microcontroller", "micro controller", "微控制器", "单片机", "soc")):
        return "MCU"
    if any(token in folded for token in ("expander", "io expander", "gpio expander", "扩展芯片", "io扩展", "gpio扩展")):
        return "扩展芯片"
    if any(token in folded for token in ("transceiver", "收发器")):
        return "收发器"
    if any(token in folded for token in ("driver", "驱动芯片", "gate driver")):
        return "驱动芯片"
    if any(token in folded for token in ("power", "regulator", "电源芯片")):
        return "电源芯片"
    if any(token in folded for token in ("connector", "连接器")):
        return "连接器"
    if any(token in folded for token in ("ic", "integrated circuit", "芯片")):
        return "其他IC"
    if raw:
        return raw
    return "未确认"


def _function_owner(device_class: str) -> str:
    if device_class == "MCU":
        return "MCU"
    if device_class == "扩展芯片":
        return "扩展芯片"
    if device_class in ("未确认", ""):
        return "未确认"
    return "其他器件"


def _style_sheet(ws, headers: list[str]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for col_idx, header in enumerate(headers, 1):
        values = [len(header)]
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            value = _text(row[0].value)
            values.append(min(len(value), 48))
            row[0].alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(max(values) + 2, 42))
    if ws.max_row >= 2 and ws.max_column:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        tab = Table(displayName="Table" + "".join(c for c in ws.title if c.isalnum())[:20], ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(tab)


def export_inventory(inventory: dict[str, Any], output_path: Path) -> int:
    components = inventory.get("components", [])
    pins = inventory.get("pins", [])
    nets = inventory.get("nets", [])
    if not isinstance(components, list) or not isinstance(pins, list) or not isinstance(nets, list):
        raise ValueError("JSON fields components, pins and nets must be arrays.")
    component_by_ref = {str(c.get("ref")): c for c in components if isinstance(c, dict) and c.get("ref") is not None}
    net_by_id = {str(n.get("net_id")): n for n in nets if isinstance(n, dict) and n.get("net_id") is not None}

    wb = Workbook()
    pin_ws = wb.active
    pin_ws.title = "PIN信号清单"
    pin_ws.append(PIN_HEADERS)
    for pin in pins:
        if not isinstance(pin, dict):
            continue
        component = component_by_ref.get(str(pin.get("component_ref")), {})
        net = net_by_id.get(str(pin.get("net_id")), {})
        device_class = _device_class(component)
        signal_name = pin.get("signal_name") or pin.get("signal") or net.get("name_as_drawn") or net.get("normalized_name")
        endpoints = net.get("endpoints", [])
        pin_ws.append([
            _text(signal_name), _text(net.get("name_as_drawn")), _text(pin.get("net_id")),
            device_class, _text(component.get("ref")), _text(component.get("part_number")),
            _text(component.get("package")), _text(pin.get("pin_number")),
            _text(pin.get("pin_name_as_drawn") or pin.get("pin_name")),
            _text(pin.get("datasheet_function") or pin.get("function")),
            _function_owner(device_class), _text(pin.get("signal_direction") or net.get("signal_direction")),
            _text(pin.get("electrical_domain") or net.get("electrical_domain")),
            _text(pin.get("connected_to") or endpoints),
            _text(pin.get("page") or pin.get("schematic_page")),
            _text(pin.get("sheet") or pin.get("schematic_sheet")),
            _text(pin.get("datasheet_reference") or pin.get("datasheet_ref")),
            _evidence(pin.get("evidence")), _text(pin.get("confidence")),
            _text(pin.get("review_state") or net.get("review_state"))
        ])
    _style_sheet(pin_ws, PIN_HEADERS)

    comp_ws = wb.create_sheet("器件清单")
    comp_ws.append(COMPONENT_HEADERS)
    for component in components:
        if not isinstance(component, dict):
            continue
        comp_ws.append([
            _text(component.get("ref")), _device_class(component), _text(component.get("part_number")),
            _text(component.get("package")), _text(component.get("value")),
            _text(component.get("population")), _text(component.get("role_hypothesis") or component.get("function")),
            _text(component.get("confidence")), _text(component.get("review_state")),
            _evidence(component.get("evidence"))
        ])
    _style_sheet(comp_ws, COMPONENT_HEADERS)

    net_ws = wb.create_sheet("网络与跨页")
    net_ws.append(NET_HEADERS)
    for net in nets:
        if not isinstance(net, dict):
            continue
        links = net.get("cross_sheet_links") or [{}]
        if not isinstance(links, list):
            links = [{}]
        for link in links:
            if not isinstance(link, dict):
                link = {}
            net_ws.append([
                _text(net.get("net_id")), _text(net.get("name_as_drawn")), _text(net.get("normalized_name")),
                _text(net.get("electrical_domain")), _text(net.get("signal_direction")),
                _text(net.get("endpoints")), _text(link.get("source_sheet")), _text(link.get("source_page")),
                _text(link.get("source_label")), _text(link.get("destination_sheet")),
                _text(link.get("destination_page")), _text(link.get("destination_label")),
                _text(link.get("status")), _text(link.get("reason") or link.get("evidence")),
                _text(net.get("confidence")), _text(net.get("review_state"))
            ])
    _style_sheet(net_ws, NET_HEADERS)

    guide = wb.create_sheet("字段说明")
    guide.append(["字段/规则", "说明"])
    guide_rows = [
        ("器件类别", "优先读取 components[].device_type/component_type/kind/category；MCU 和扩展芯片会单独分类。未知类型不会自动猜测。"),
        ("功能归属", "MCU、扩展芯片或其他器件；这是根据清单中的器件类别归类，不代表已证明软件运行时的数据流。"),
        ("PIN信号清单", "每个 pins 数组记录生成一行；器件/网络信息按 component_ref 和 net_id 关联。"),
        ("证据与置信度", "请填入原理图页码、位号、引脚、数据手册版本/页码及来源。空白字段表示输入数据没有提供，不会自动补造。"),
        ("示例数据", "示例硬件清单为占位数据，导出结果不可当作真实 ECU 引脚表。"),
        ("限制", "Excel 导出只整理已有 JSON 数据，不执行 OCR、原理图识图、电气连通性重建或 AUTOSAR 语义验证。")
    ]
    for row in guide_rows:
        guide.append(row)
    _style_sheet(guide, ["字段/规则", "说明"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    # Reopen the produced workbook so a malformed/unreadable output fails the command.
    check = load_workbook(output_path, read_only=True, data_only=True)
    try:
        if "PIN信号清单" not in check.sheetnames:
            raise RuntimeError("Generated workbook is missing the PIN signal sheet.")
        count = max(check["PIN信号清单"].max_row - 1, 0)
    finally:
        check.close()
    return count


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inventory_json", type=Path, help="Hardware inventory JSON input")
    parser.add_argument("--out", type=Path, default=Path("engineering/reports/pin_signal_inventory.xlsx"))
    args = parser.parse_args()
    if not args.inventory_json.is_file():
        parser.error(f"Input JSON not found: {args.inventory_json}")
    try:
        inventory = json.loads(args.inventory_json.read_text(encoding="utf-8"))
        if not isinstance(inventory, dict):
            raise ValueError("Inventory root must be a JSON object.")
        count = export_inventory(inventory, args.out)
    except (OSError, json.JSONDecodeError, ValueError, RuntimeError) as exc:
        print(f"Export failed: {exc}")
        return 2
    print(f"Workbook: {args.out}")
    print(f"PIN rows exported: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
